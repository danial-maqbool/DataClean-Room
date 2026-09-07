"""Native clean copies remove source objects rather than hiding them."""

import base64
import io
import json
import shutil
import subprocess
import zipfile
from unittest.mock import patch
from app.native import request, clean_office
from app.sanitizer import analyze
from localdesk.safety import InputError, digest
from tests.fixtures import make_pdf, make_image
from tests.support import AppCase


class NativePrivacyTests(AppCase):
    def scan_and_clean(self, path, **settings):
        before = digest(path)
        scan = self.finish("scan", {"paths": [str(path)], **settings})
        ident = scan["scans"][0]["id"]
        result = self.finish("sanitize", {"id": ident})
        self.assertEqual(digest(path), before)
        return self.output(result["output"]), result, ident

    def test_pdf_copy_has_new_pixels_and_no_source_text(self):
        from pypdf import PdfReader

        p = make_pdf(self.workspace / "source.pdf", text="Email alex@example.test")
        raw, result, _ = self.scan_and_clean(p)
        pdf = PdfReader(io.BytesIO(raw), strict=True)
        self.assertEqual(len(pdf.pages), 1)
        self.assertEqual(pdf.pages[0].extract_text(), "")
        self.assertNotIn(b"alex@example.test", raw)
        self.assertFalse(result["audit"]["source_objects_copied"])
        self.assertGreater(result["audit"]["painted_regions"], 0)

    def test_pdf_preview_and_manual_regions(self):
        from PIL import Image

        p = make_pdf(self.workspace / "page.pdf", text="Public text")
        scan = self.finish("scan", {"paths": [str(p)]})
        ident = scan["scans"][0]["id"]
        image = self.app.get("image", {"id": ident, "page": "1"})
        with Image.open(io.BytesIO(base64.b64decode(image["data"]))) as im:
            self.assertEqual(im.size, (800, 600))
        result = self.finish(
            "sanitize",
            {"id": ident, "page_rectangles": [{"page": 1, "rect": [0, 0, 80, 80]}]},
        )
        self.assertEqual(result["audit"]["selected_regions"], 1)
        self.assertTrue(self.output(result["output"]).startswith(b"%PDF"))

    def test_pdf_outside_region_fails(self):
        p = make_pdf(self.workspace / "page.pdf")
        with self.assertRaisesRegex(InputError, "outside"):
            request(
                p.read_bytes(),
                ".pdf",
                "pdf",
                rectangles=[{"page": 1, "rect": [0, 0, 99999, 500]}],
            )

    def test_pdf_source_attachments_do_not_survive(self):
        from pypdf import PdfWriter, PdfReader

        p = make_pdf(self.workspace / "page.pdf")
        writer = PdfWriter()
        writer.append(str(p))
        writer.add_attachment("private.txt", b"UNIQUE_ATTACHMENT_SECRET")
        writer.add_metadata({"/Author": "UNIQUE_PRIVATE_AUTHOR"})
        stream = io.BytesIO()
        writer.write(stream)
        p.write_bytes(stream.getvalue())
        raw, _, _ = self.scan_and_clean(p)
        pdf = PdfReader(io.BytesIO(raw))
        self.assertFalse(pdf.attachments)
        self.assertNotIn(b"UNIQUE_ATTACHMENT_SECRET", raw)
        self.assertNotIn("UNIQUE_PRIVATE_AUTHOR", str(pdf.metadata))

    def test_docx_reconstruction_removes_values_and_source_metadata(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Email alex@example.test")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Public total 42"
        document.core_properties.author = "UNIQUE_PRIVATE_AUTHOR"
        p = self.workspace / "report.docx"
        document.save(p)
        raw, result, _ = self.scan_and_clean(p)
        doc = Document(io.BytesIO(raw))
        self.assertNotIn("alex@example.test", doc.paragraphs[0].text)
        self.assertEqual(doc.tables[0].cell(0, 0).text, "Public total 42")
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            text = "".join(z.read(i).decode("utf-8", "ignore") for i in z.namelist())
            self.assertNotIn("UNIQUE_PRIVATE_AUTHOR", text)
            self.assertNotIn("alex@example.test", text)
        self.assertEqual(result["audit"]["output_format"], ".docx")

    def test_word_hidden_runs_and_headers_are_omitted(self):
        from docx import Document

        document = Document()
        p = document.add_paragraph("Public text. ")
        hidden = p.add_run("UNIQUE_HIDDEN")
        hidden.font.hidden = True
        document.sections[0].header.paragraphs[0].text = "PRIVATE_HEADER"
        stream = io.BytesIO()
        document.save(stream)
        raw, _ = clean_office(stream.getvalue(), ".docx")
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            xml = "".join(z.read(i).decode("utf-8", "ignore") for i in z.namelist())
            self.assertNotIn("UNIQUE_HIDDEN", xml)
            self.assertNotIn("PRIVATE_HEADER", xml)

    def test_xlsx_hidden_sheets_formulas_and_comments_removed(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.comments import Comment

        book = Workbook()
        sheet = book.active
        sheet["A1"] = "alex@example.test"
        sheet["B1"] = 42
        sheet["A2"] = "=1+1"
        sheet["B1"].comment = Comment("UNIQUE_PRIVATE_COMMENT", "Private author")
        hidden = book.create_sheet("PRIVATE_SHEET")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "UNIQUE_HIDDEN"
        book.properties.creator = "UNIQUE_AUTHOR"
        p = self.workspace / "book.xlsx"
        book.save(p)
        book.close()
        raw, result, _ = self.scan_and_clean(p)
        out = load_workbook(io.BytesIO(raw))
        self.assertEqual(out.sheetnames, ["Sheet 1"])
        self.assertNotIn("alex@example.test", out.active["A1"].value)
        self.assertEqual(out.active["B1"].value, 42)
        self.assertEqual(out.active["A2"].value, "[FORMULA REMOVED]")
        self.assertIsNone(out.active["B1"].comment)
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            xml = "".join(z.read(i).decode("utf-8", "ignore") for i in z.namelist())
            for value in [
                "UNIQUE_PRIVATE_COMMENT",
                "UNIQUE_HIDDEN",
                "PRIVATE_SHEET",
                "UNIQUE_AUTHOR",
            ]:
                self.assertNotIn(value, xml)
        out.close()

    def test_pptx_rebuilds_redacted_slide_text(self):
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = (
            "Contact alex@example.test"
        )
        deck.core_properties.author = "UNIQUE_PRIVATE_AUTHOR"
        p = self.workspace / "slides.pptx"
        deck.save(p)
        raw, result, _ = self.scan_and_clean(p)
        out = Presentation(io.BytesIO(raw))
        self.assertEqual(len(out.slides), 1)
        text = " ".join(
            shape.text for shape in out.slides[0].shapes if shape.has_text_frame
        )
        self.assertIn("[EMAIL_1]", text)
        self.assertNotIn("alex@example.test", text)
        self.assertFalse(out.core_properties.author)

    def test_actual_ocr_image_removes_detected_pixels(self):
        if not shutil.which("tesseract"):
            self.skipTest("Tesseract is required for this integration test.")
        from PIL import Image

        p = make_image(
            self.workspace / "scan.png", text="PRIVATEVALUE\nPublic invoice 42"
        )
        raw, result, _ = self.scan_and_clean(p, ocr=True, literals=["PRIVATEVALUE"])
        self.assertGreater(result["audit"]["automatic_regions"], 0)
        with Image.open(io.BytesIO(raw)) as im:
            self.assertEqual(im.info, {})

    def test_native_worker_rejects_timeout(self):
        with patch(
            "app.native.subprocess.run",
            side_effect=subprocess.TimeoutExpired("test", 150),
        ):
            with self.assertRaisesRegex(InputError, "150 second"):
                request(b"%PDF-", ".pdf", "layout")

    def test_native_worker_preserves_error_message(self):
        result = subprocess.CompletedProcess(
            [], 0, json.dumps({"error": "The file needs a password."}), ""
        )
        with patch("app.native.subprocess.run", return_value=result):
            with self.assertRaisesRegex(InputError, "needs a password"):
                request(b"%PDF-", ".pdf", "layout")

    def test_office_entities_rejected_before_library_parsing(self):
        from tests.support import zipped

        payload = zipped(
            {
                "word/document.xml": '<!DOCTYPE d [<!ENTITY x SYSTEM "file:///etc/passwd">]><d>&x;</d>'
            }
        )
        with self.assertRaises(InputError):
            clean_office(payload, ".docx")

    def test_xlsx_grouped_hidden_columns_and_merged_cells(self):
        from openpyxl import Workbook, load_workbook

        book = Workbook()
        sheet = book.active
        sheet["A1"] = "Public"
        sheet["B1"] = "PRIVATE_B"
        sheet["C1"] = "PRIVATE_C"
        sheet.column_dimensions.group("B", "C", hidden=True)
        sheet.merge_cells("A2:C2")
        sheet["A2"] = "Merged public"
        buf = io.BytesIO()
        book.save(buf)
        book.close()
        raw, _ = clean_office(buf.getvalue(), ".xlsx")
        result = load_workbook(io.BytesIO(raw))
        values = str(list(result.active.values))
        result.close()
        self.assertNotIn("PRIVATE_B", values)
        self.assertNotIn("PRIVATE_C", values)
        self.assertIn("Merged public", values)

    def test_pdf_rotated_crop_preserves_visible_size_and_masks(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import RectangleObject
        from PIL import Image

        for rotation in (0, 90, 180, 270):
            p = make_pdf(
                self.workspace / f"cropped-{rotation}.pdf",
                text="Email alex@example.test",
            )
            reader = PdfReader(str(p))
            page = reader.pages[0]
            page.cropbox = RectangleObject([20, 20, 380, 280])
            page.rotate(rotation)
            writer = PdfWriter()
            writer.add_page(page)
            buf = io.BytesIO()
            writer.write(buf)
            raw = buf.getvalue()
            preview = request(raw, ".pdf", "preview")
            expected = (720, 520) if rotation in (0, 180) else (520, 720)
            self.assertEqual((preview["width"], preview["height"]), expected)
            result = request(raw, ".pdf", "pdf", literals=["Email"])
            self.assertGreater(result["audit"]["painted_regions"], 0)
            output = PdfReader(io.BytesIO(result["content"]))
            self.assertEqual(output.pages[0].extract_text(), "")
            with Image.open(io.BytesIO(preview["content"])) as image:
                self.assertEqual(image.size, expected)
