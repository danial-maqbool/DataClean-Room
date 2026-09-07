"""Synthetic acceptance operations using each repository's own installed runtime."""

from pathlib import Path
import csv
import hashlib
import secrets
import io
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import time
import traceback
import zipfile

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))
from app.service import Application
from localdesk.safety import InputError
from tests.fixtures import make_image, make_pdf

import argparse

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument(
    "--report-dir",
    type=Path,
    required=True,
    help="New ignored evidence folder for this run. Existing folders are preserved.",
)
args = parser.parse_args()
BASE = args.report_dir.resolve()
BASE.mkdir(parents=True, exist_ok=False)
WORK = BASE / "synthetic space unicode-ÃƒÂ©"
WORK.mkdir()
RESULTS = []
APP = None
PASSWORD = secrets.token_urlsafe(32)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def record(name, func):
    started = time.monotonic()
    try:
        detail = func()
        result = dict(check=name, status="PASS", detail=detail)
    except Exception:
        result = dict(check=name, status="FAIL", traceback=traceback.format_exc())
    result["seconds"] = time.monotonic() - started
    RESULTS.append(result)
    (BASE / "results.json").write_text(
        json.dumps(RESULTS, indent=2, default=str), encoding="utf-8"
    )
    print(name, result["status"], flush=True)


def finish(action, body, expected="done"):
    ident = APP.common_post(action, body)["job_id"]
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        job = APP.jobs.get(ident)
        if job["status"] not in ("running", "queued"):
            assert job["status"] == expected, job
            return job["result"] if expected == "done" else job
        time.sleep(0.03)
    raise TimeoutError(action)


def output(artifact):
    return APP.download(artifact["path"]).read_bytes()


def restart():
    global APP
    APP.close()
    APP = Application(
        ROOT,
        BASE / "persistent-data",
        passphrase=PASSWORD,
    )


def file(name, content):
    path = WORK / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content.encode("utf-8") if isinstance(content, str) else content)
    return path


def office_fixtures():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    doc = Document()
    doc.add_paragraph("Office canary qa@example.test public number 42")
    doc.core_properties.author = "PRIVATE_AUTHOR_CANARY"
    doc.save(WORK / "office.docx")
    book = Workbook()
    book.active.append(["Office canary qa@example.test", 42])
    book.save(WORK / "office.xlsx")
    book.close()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = (
        "Office canary qa@example.test public number 42"
    )
    deck.save(WORK / "office.pptx")
    return [WORK / ("office" + ext) for ext in (".docx", ".xlsx", ".pptx")]


def privacy():
    def office():
        evidence = []
        for path in office_fixtures():
            before = sha(path)
            scan = finish("scan", {"paths": [str(path)]})
            clean = finish("sanitize", {"id": scan["scans"][0]["id"]})
            raw = output(clean["output"])
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                parts = b"\n".join(z.read(n) for n in z.namelist())
                assert b"qa@example.test" not in parts
                assert b"PRIVATE_AUTHOR_CANARY" not in parts
                assert b"42" in parts
            assert sha(path) == before
            evidence.append(clean)
        return evidence

    record("DC-01/DC-05/DC-08 Office exported package canaries and hashes", office)

    def images():
        from PIL import Image

        p = WORK / "gps.jpg"
        im = Image.new("RGB", (160, 100), "white")
        exif = Image.Exif()
        exif[270] = "PRIVATE_EXIF_CANARY"
        exif[34853] = {1: "N", 2: (31.0, 30.0, 0.0), 3: "E", 4: (74.0, 20.0, 0.0)}
        im.save(p, exif=exif)
        with Image.open(p) as source:
            assert source.getexif().get_ifd(34853)
        before = sha(p)
        scan = finish("scan", {"paths": [str(p)]})
        clean = finish(
            "sanitize", {"id": scan["scans"][0]["id"], "rectangles": [[10, 10, 90, 70]]}
        )
        raw = output(clean["output"])
        with Image.open(io.BytesIO(raw)) as out:
            assert not out.getexif()
            assert all(
                out.convert("RGB").getpixel((x, y)) == (0, 0, 0)
                for x in range(12, 88)
                for y in range(12, 68)
            )
        assert b"PRIVATE_EXIF_CANARY" not in raw
        assert sha(p) == before
        return clean

    record("DC-03 Independent full-region exported pixels and EXIF GPS", images)

    def pdfs():
        import pypdfium2 as pdfium
        from pypdf import PdfReader, PdfWriter

        paths = [make_pdf(WORK / "native.pdf", text="QA_SECRET_CANARY public 42")]
        image = make_image(WORK / "scan.png", text="QA_SECRET_CANARY\nPublic number 42")
        from PIL import Image

        with Image.open(image) as im:
            im.convert("RGB").save(WORK / "scanned.pdf", "PDF")
        paths.append(WORK / "scanned.pdf")
        writer = PdfWriter()
        for path in paths:
            writer.append(str(path))
        writer.pages[0].rotate(90)
        writer.add_attachment("canary.txt", b"QA_SECRET_ATTACHMENT")
        writer.add_metadata({"/Author": "QA_SECRET_AUTHOR"})
        writer.write(WORK / "mixed-rotated.pdf")
        paths.append(WORK / "mixed-rotated.pdf")
        results = []
        for path in paths:
            before = sha(path)
            scan = finish(
                "scan",
                {"paths": [str(path)], "ocr": True, "literals": ["QA_SECRET_CANARY"]},
            )
            clean = finish("sanitize", {"id": scan["scans"][0]["id"]})
            raw = output(clean["output"])
            reader = PdfReader(io.BytesIO(raw))
            assert not reader.attachments
            assert not reader.get_fields()
            assert "QA_SECRET" not in str(reader.metadata)
            assert all(
                not p.extract_text().strip() and not p.get("/Annots")
                for p in reader.pages
            )
            assert b"QA_SECRET" not in raw
            with pdfium.PdfDocument(raw) as pdf:
                for number in range(len(pdf)):
                    page = pdf[number]
                    bitmap = page.render(scale=1)
                    rendered = bitmap.to_pil()
                    target = BASE / f"{path.stem}-clean-page-{number}.png"
                    rendered.save(target)
                    ocr = subprocess.run(
                        ["tesseract", str(target), "stdout", "-l", "eng"],
                        capture_output=True,
                        text=True,
                        check=True,
                    ).stdout
                    assert "QA_SECRET_CANARY" not in ocr, ocr
                    (target.with_suffix(".ocr.txt")).write_text(ocr)
                    rendered.close()
                    bitmap.close()
                    page.close()
            assert sha(path) == before
            results.append(clean)
        return results

    record(
        "DC-04/DC-08 native scanned mixed rotated PDF independent render and OCR", pdfs
    )

    def reviewed_pdf():
        import base64
        import pypdfium2 as pdfium
        from PIL import Image
        from pypdf import PdfReader

        details = []
        for path in (WORK / "scanned.pdf", WORK / "mixed-rotated.pdf"):
            scan = finish(
                "scan",
                {"paths": [str(path)], "ocr": True, "literals": ["QA_SECRET_CANARY"]},
            )
            ident = scan["scans"][0]["id"]
            pages = len(PdfReader(str(path)).pages)
            regions = []
            for number in range(1, pages + 1):
                preview = APP.get("image", {"id": ident, "page": str(number)})
                with Image.open(io.BytesIO(base64.b64decode(preview["data"]))) as im:
                    width, height = im.size
                # Withhold the rotated native page in full; on scans select only
                # the known first-line canary band, retaining the public line.
                bottom = (
                    height
                    if path.name == "mixed-rotated.pdf" and number == 1
                    else int(height * 0.4)
                )
                regions.append({"page": number, "rect": [0, 0, width, bottom]})
            clean = finish("sanitize", {"id": ident, "page_rectangles": regions})
            raw = output(clean["output"])
            with pdfium.PdfDocument(raw) as pdf:
                for number in range(len(pdf)):
                    page = pdf[number]
                    bitmap = page.render(scale=2)
                    im = bitmap.to_pil()
                    target = BASE / f"{path.stem}-reviewed-{number}.png"
                    im.save(target)
                    text = subprocess.check_output(
                        ["tesseract", str(target), "stdout", "-l", "eng"],
                        text=True,
                        stderr=subprocess.DEVNULL,
                    )
                    assert "SECRET_CANARY" not in text, text
                    if regions[number]["rect"][3] < im.height:
                        assert "42" in text, text
                    im.close()
                    bitmap.close()
                    page.close()
            details.append({"manual_regions": regions, "result": clean})
        return details

    record(
        "DC-04 reviewed manual masks remove missed OCR canary from exported pixels",
        reviewed_pdf,
    )

    def changed():
        results = []
        for mode in ("same-size", "different-size", "removed"):
            path = file(mode + ".txt", "qa@example.test")
            scan = finish("scan", {"paths": [str(path)]})
            if mode == "removed":
                path.unlink()
            else:
                path.write_text(
                    "ab@example.test" if mode == "same-size" else "Changed source size"
                )
            try:
                result = APP.common_post("sanitize", {"id": scan["scans"][0]["id"]})
            except InputError as exc:
                results.append(str(exc))
                continue
            deadline = time.monotonic() + 30
            while time.monotonic() < deadline:
                job = APP.jobs.get(result["job_id"])
                if job["status"] not in ("queued", "running"):
                    break
                time.sleep(0.02)
            assert job["status"] == "failed", job
            results.append(job["error"])
        return results

    record("DC-07 same-size size-change and removed source refusal", changed)


def shared():
    def persist():
        APP.store.set("qa_canary", "PERSISTED_SYNTHETIC_CANARY")
        backup = APP.common_post("backup", {})
        raw = output(backup)
        assert b"PERSISTED_SYNTHETIC_CANARY" not in raw
        restart()
        assert APP.store.get("qa_canary") == "PERSISTED_SYNTHETIC_CANARY"
        from localdesk.storage import Store

        restored = BASE / "restored.vault"
        restored.write_bytes(raw)
        store = Store(restored, PASSWORD)
        assert store.get("qa_canary") == "PERSISTED_SYNTHETIC_CANARY"
        store.close()
        before = sha(restored)
        try:
            Store(restored, secrets.token_urlsafe(32))
        except InputError:
            pass
        else:
            raise AssertionError("Wrong password accepted")
        assert sha(restored) == before
        return {
            "backup_sha256": hashlib.sha256(raw).hexdigest(),
            "restart_and_restore": True,
        }

    record(
        "Shared real encrypted persistence backup restore wrong-password preservation",
        persist,
    )


def main():
    global APP
    APP = Application(
        ROOT,
        BASE / "persistent-data",
        passphrase=PASSWORD,
    )
    try:
        privacy()
        shared()
    finally:
        APP.close()
        hashes = {
            p.relative_to(BASE).as_posix(): sha(p)
            for p in BASE.rglob("*")
            if p.is_file() and p.name != "hashes.json"
        }
        (BASE / "hashes.json").write_text(
            json.dumps(hashes, indent=2), encoding="utf-8"
        )
    return int(any(r["status"] == "FAIL" for r in RESULTS))


if __name__ == "__main__":
    raise SystemExit(main())
