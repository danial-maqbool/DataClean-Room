"""Copy-only PDF and Office reconstruction with explicit omissions.

PDF output contains new page pixels, not original PDF objects. Office output
contains rebuilt text and cells. No source macros, comments, links, or media
objects are copied. Detection remains a review aid, not a privacy guarantee.
"""

from __future__ import annotations

import base64
from collections import Counter
import hashlib
import io
import json
import math
from pathlib import Path
import subprocess
import sys
import tempfile
from xml.etree import ElementTree as ET

from localdesk.safety import InputError, MAX_FILE_BYTES, open_zip, safe_xml, integer
from .detectors import detect, public_matches, redact
from .ocr_matching import document_matches, review_warnings, require_review

OFFICE = {".docx", ".xlsx", ".pptx"}
MAX_PIXELS = 20_000_000


def request(raw, suffix, action, **options):
    if len(raw) > MAX_FILE_BYTES:
        raise InputError("This source exceeds 25 MiB.")
    envelope = dict(
        raw=base64.b64encode(raw).decode("ascii"),
        suffix=suffix,
        action=action,
        options=options,
    )
    try:
        result = subprocess.run(
            [sys.executable, "-m", "app.privacy_worker"],
            input=json.dumps(envelope),
            text=True,
            encoding="utf-8",
            capture_output=True,
            timeout=150,
            cwd=Path(__file__).resolve().parents[1],
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise InputError(
            "Privacy reconstruction reached its 150 second limit. Split this file."
        ) from exc
    if result.returncode or len(result.stdout) > 38_000_000:
        raise InputError("The privacy worker stopped without a bounded result.")
    try:
        response = json.loads(result.stdout)
        if response.get("error"):
            raise InputError(response["error"])
        if "content" in response:
            response["content"] = base64.b64decode(response["content"], validate=True)
            if len(response["content"]) > MAX_FILE_BYTES:
                raise InputError("The clean output exceeds 25 MiB.")
        return response
    except InputError:
        raise
    except (ValueError, TypeError, AttributeError) as exc:
        raise InputError("The privacy worker returned an invalid result.") from exc


def _validate_office(raw):
    with open_zip(raw) as archive:
        for item in archive.infolist():
            if item.filename.endswith((".xml", ".rels")):
                safe_xml(archive.read(item))


def _clear_properties(properties):
    for key in (
        "author",
        "last_modified_by",
        "title",
        "subject",
        "keywords",
        "comments",
        "category",
        "identifier",
        "language",
        "version",
        "content_status",
    ):
        if hasattr(properties, key):
            try:
                setattr(properties, key, "")
            except (ValueError, TypeError, AttributeError):
                pass


def clean_office(raw, suffix, *, enabled=None, literals=None, mode="labels"):
    _validate_office(raw)
    count = 0
    characters = 0

    def clean(value):
        nonlocal count, characters
        text = str(value)
        characters += len(text)
        if characters > 250_000:
            raise InputError(
                "Office reconstruction exceeds 250,000 text characters. Split the source."
            )
        matches = detect(text, enabled=enabled, literals=literals)
        count += len(matches)
        return redact(text, matches, mode)

    output = io.BytesIO()
    omissions = []
    if suffix == ".docx":
        from docx import Document

        fresh = Document()
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with open_zip(raw) as source:
            document = safe_xml(source.read("word/document.xml"))
            hidden_styles = set()
            style_parents = {}
            if "word/styles.xml" in source.namelist():
                for style in safe_xml(source.read("word/styles.xml")).iter(W + "style"):
                    ident = style.attrib.get(W + "styleId")
                    parent = style.find("./" + W + "basedOn")
                    if parent is not None:
                        style_parents[ident] = parent.attrib.get(W + "val")
                    if (
                        style.find(".//" + W + "vanish") is not None
                        or style.find(".//" + W + "webHidden") is not None
                    ):
                        hidden_styles.add(ident)
                for _ in range(len(style_parents)):
                    added = {
                        child
                        for child, parent in style_parents.items()
                        if parent in hidden_styles
                    } - hidden_styles
                    if not added:
                        break
                    hidden_styles.update(added)

            def visible(node):
                local = node.tag.rsplit("}", 1)[-1]
                if local in {
                    "del",
                    "moveFrom",
                    "instrText",
                    "delText",
                    "drawing",
                    "pict",
                    "object",
                }:
                    return ""
                if local == "r":
                    if (
                        node.find("./" + W + "rPr/" + W + "vanish") is not None
                        or node.find("./" + W + "rPr/" + W + "webHidden") is not None
                    ):
                        return ""
                    style = node.find("./" + W + "rPr/" + W + "rStyle")
                    if (
                        style is not None
                        and style.attrib.get(W + "val") in hidden_styles
                    ):
                        return ""
                if local == "p":
                    style = node.find("./" + W + "pPr/" + W + "pStyle")
                    if (
                        style is not None
                        and style.attrib.get(W + "val") in hidden_styles
                    ):
                        return ""
                if local == "t":
                    return node.text or ""
                if local in {"br", "cr"}:
                    return "\n"
                if local == "tab":
                    return "\t"
                return "".join(visible(child) for child in node)

            body = document.find(W + "body")
            if body is None:
                raise InputError("The Word document has no readable body.")
            blocks = list(body)
            if len(blocks) > 5000:
                raise InputError("Word reconstruction exceeds 5,000 body blocks.")
            for block in blocks:
                if block.tag == W + "p":
                    text = visible(block)
                    if text:
                        fresh.add_paragraph(clean(text))
                elif block.tag == W + "tbl":
                    rows = block.findall(W + "tr")
                    width = max((len(row.findall(W + "tc")) for row in rows), default=0)
                    if len(rows) > 1000 or width > 100:
                        raise InputError(
                            "The Word table exceeds reconstruction limits."
                        )
                    if width:
                        table = fresh.add_table(rows=len(rows), cols=width)
                        for r, row in enumerate(rows):
                            for c, cell in enumerate(row.findall(W + "tc")):
                                table.cell(r, c).text = clean(
                                    "\n".join(visible(p) for p in cell.findall(W + "p"))
                                )
        _clear_properties(fresh.core_properties)
        fresh.save(output)
        # Parse the new document before returning it.
        Document(io.BytesIO(output.getvalue()))
        omissions = [
            "Headers, footers, embedded objects, pictures, comments, tracked deletions, and hidden runs were not copied.",
            "Body paragraphs and table cells were rebuilt. Source styling, fields, and hyperlinks were not retained.",
        ]
    elif suffix == ".xlsx":
        from openpyxl import Workbook, load_workbook
        from openpyxl.utils import get_column_letter, column_index_from_string

        source = load_workbook(
            io.BytesIO(raw), data_only=False, read_only=False, keep_links=False
        )
        fresh = Workbook()
        fresh.remove(fresh.active)
        cells = 0
        try:
            for sheet in source.worksheets:
                if sheet.sheet_state != "visible":
                    continue
                if sheet.max_row > 50_000 or sheet.max_column > 200:
                    raise InputError(
                        "The worksheet exceeds 50,000 rows or 200 columns."
                    )
                target = fresh.create_sheet("Sheet " + str(len(fresh.worksheets) + 1))
                hidden_columns = set()
                for label, dimension in sheet.column_dimensions.items():
                    if dimension.hidden:
                        first = dimension.min or column_index_from_string(label)
                        last = dimension.max or first
                        hidden_columns.update(range(first, min(last, 200) + 1))
                for row in sheet.iter_rows():
                    if sheet.row_dimensions[row[0].row].hidden:
                        continue
                    values = []
                    for cell in row:
                        if cell.column in hidden_columns:
                            continue
                        cells += 1
                        if cells > 100_000:
                            raise InputError("The workbook exceeds 100,000 cells.")
                        value = cell.value
                        if cell.data_type == "f":
                            value = "[FORMULA REMOVED]"
                        elif value is not None:
                            rewritten = clean(value)
                            if rewritten != str(value) or isinstance(value, str):
                                value = rewritten
                            if isinstance(value, str) and value.lstrip().startswith(
                                ("=", "+", "-", "@")
                            ):
                                value = "'" + value
                        values.append(value)
                    target.append(values)
            if not fresh.worksheets:
                raise InputError("The workbook contains no visible worksheets.")
            fresh.properties.creator = ""
            fresh.properties.lastModifiedBy = ""
            fresh.properties.title = ""
            fresh.properties.subject = ""
            fresh.properties.description = ""
            fresh.save(output)
            check = load_workbook(
                io.BytesIO(output.getvalue()), read_only=True, data_only=True
            )
            check.close()
        finally:
            source.close()
            fresh.close()
        omissions = [
            "Hidden worksheets, hidden rows and columns, formulas, links, comments, pictures, charts, and macros were not copied.",
            "Visible cells were rebuilt with generic worksheet names. Original formatting and calculations were not retained.",
        ]
    elif suffix == ".pptx":
        from pptx import Presentation
        from pptx.util import Pt

        source = Presentation(io.BytesIO(raw))
        fresh = Presentation()
        fresh.slide_width, fresh.slide_height = source.slide_width, source.slide_height
        if not 1 <= len(source.slides) <= 100:
            raise InputError("Use a presentation with 1 to 100 slides.")
        for old_slide in source.slides:
            if old_slide._element.get("show") == "0":
                continue
            slide = fresh.slides.add_slide(fresh.slide_layouts[6])
            for shape in old_slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = "\n".join(p.text for p in shape.text_frame.paragraphs)
                elif getattr(shape, "has_table", False):
                    text = "\n".join(
                        " | ".join(cell.text for cell in row.cells)
                        for row in shape.table.rows
                    )
                else:
                    continue
                if not text:
                    continue
                target = slide.shapes.add_textbox(
                    shape.left,
                    shape.top,
                    max(shape.width, 91440),
                    max(shape.height, 91440),
                )
                target.text_frame.word_wrap = True
                target.text_frame.text = clean(text)
                for paragraph in target.text_frame.paragraphs:
                    paragraph.font.size = Pt(14)
        if not fresh.slides:
            raise InputError("The presentation contains no visible slides.")
        _clear_properties(fresh.core_properties)
        fresh.save(output)
        Presentation(io.BytesIO(output.getvalue()))
        omissions = [
            "Hidden slides, speaker notes, links, comments, pictures, charts, embedded files, and macros were not copied.",
            "Slide text and table text were rebuilt in their original positions. Source styles, transitions, and graphics were not retained.",
        ]
    else:
        raise InputError("Native Office reconstruction supports DOCX, XLSX, and PPTX.")
    data = output.getvalue()
    if len(data) > MAX_FILE_BYTES:
        raise InputError("The clean Office copy exceeds 25 MiB.")
    return data, {
        "removed_matches": count,
        "source_metadata_copied": False,
        "output_format": suffix,
        "reconstruction": "visible text and cells",
        "warnings": omissions,
        "human_review_required": True,
        "visible_content_fully_verified": False,
    }


def _layout(path, ocr):
    from localdesk.document_worker import layout

    result = layout(path, ocr=ocr, max_pages=50)
    if result.get("truncated"):
        raise InputError(
            "The document layout is incomplete. Split the input before redaction."
        )
    if path.suffix.lower() == ".pdf":
        # Rotated text may sort into columns of single letters. Also inspect
        # native character order, with exact character boxes, when it differs
        # from the positional word order. Keep the original scan and OCR.
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            for source, page in zip(pdf.pages, result["pages"]):
                chars = [
                    c
                    for c in source.chars
                    if c.get("text")
                    and c["x1"] > 0
                    and c["x0"] < page["width"]
                    and c["bottom"] > 0
                    and c["top"] < page["height"]
                ]
                if len(chars) > 250_000:
                    raise InputError(
                        "This PDF has too many native characters. Split the input."
                    )
                flow = "".join(c["text"] for c in chars)
                positional = " ".join(w["text"] for w in page["words"])
                if flow and " ".join(flow.split()) != " ".join(positional.split()):
                    if len(result["text"]) + len(flow) + 2 > 250_000:
                        raise InputError(
                            "The combined PDF text exceeds 250,000 characters. Split the input."
                        )
                    result["text"] += "\n"
                    position = len(result["text"])
                    for char in chars:
                        value = char["text"]
                        page["words"].append(
                            {
                                "text": value,
                                "start": position,
                                "end": position + len(value),
                                "bbox": [
                                    char["x0"],
                                    char["top"],
                                    char["x1"],
                                    char["bottom"],
                                ],
                                "line": "native-content-order",
                            }
                        )
                        position += len(value)
                    result["text"] += flow + "\n"
                    result["warnings"].append(
                        "Page "
                        + str(page["page"])
                        + ": also checked native character order for rotated or reordered text."
                    )
                source.close()
    return result


def _regions(words, matches, scale, width, height):
    regions = []
    for word in words:
        if any(word["start"] < m["end"] and m["start"] < word["end"] for m in matches):
            x0, y0, x1, y1 = word["bbox"]
            region = [
                max(0, math.floor(x0 * scale) - 2),
                max(0, math.floor(y0 * scale) - 2),
                min(width, math.ceil(x1 * scale) + 2),
                min(height, math.ceil(y1 * scale) + 2),
            ]
            if region[0] < region[2] and region[1] < region[3]:
                regions.append(region)
    return regions


def _undo_rotation(box, angle, width, height):
    x0, y0, x1, y1 = box
    if angle == 90:
        return [width - y1, x0, width - y0, x1]
    if angle == 180:
        return [width - x1, height - y1, width - x0, height - y0]
    if angle == 270:
        return [y0, height - x1, y1, height - x0]
    return list(box)


def _recover_orientation(path, document, *, enabled=None, literals=None):
    """Try bounded quarter turns only on pages without a custom-value candidate."""
    if not literals or (enabled is not None and "custom" not in enabled):
        return document
    from PIL import Image, ImageOps
    from localdesk.document_worker import ocr_words, render_page
    from .ocr_matching import normalized, MIN_FUZZY_LENGTH

    if not any(len(normalized(s)) >= MIN_FUZZY_LENGTH for s in literals):
        return document

    def found(page, text):
        matches, _ = document_matches(
            {"text": text, "pages": [page]}, enabled=["custom"], literals=literals
        )
        return any(
            m["start"] < w["end"] and w["start"] < m["end"]
            for m in matches
            for w in page["words"]
        )

    attempts = 0
    unresolved = []
    for page in document["pages"]:
        if found(page, document["text"]):
            continue
        if attempts >= 6:
            unresolved.append(page["page"])
            continue
        if path.suffix.lower() == ".pdf":
            image, _, _ = render_page(path, page["page"] - 1, scale=2)
            scale = 2
        else:
            with Image.open(path) as original:
                image = ImageOps.exif_transpose(original).convert("RGB")
            scale = 1
        recovered = False
        try:
            for angle in (90, 180, 270):
                if attempts >= 6:
                    break
                attempts += 1
                rotated = image.rotate(angle, expand=True)
                try:
                    with tempfile.TemporaryDirectory(
                        prefix="dataclean-orientation-"
                    ) as folder:
                        rotated_path = Path(folder) / "page.png"
                        rotated.save(rotated_path)
                        words = ocr_words(rotated_path)
                finally:
                    rotated.close()
                text = document["text"] + "\n"
                previous = None
                for word in words:
                    box = _undo_rotation(word["bbox"], angle, image.width, image.height)
                    word["match_bbox"] = [v / scale for v in word["bbox"]]
                    word["bbox"] = [v / scale for v in box]
                    if previous is not None:
                        text += " " if word["line"] == previous else "\n"
                    word["start"] = len(text)
                    text += word["text"]
                    word["end"] = len(text)
                    word["orientation"] = angle
                    previous = word["line"]
                    if len(text) > 250_000:
                        raise InputError(
                            "Orientation OCR exceeds 250,000 characters. Split the input."
                        )
                trial = {**page, "words": words}
                if found(trial, text):
                    document["text"] = text + "\n"
                    page["words"].extend(words)
                    document["warnings"].append(
                        f"Page {page['page']}: added OCR from a {angle}-degree rotation. Mask positions use the original page coordinates."
                    )
                    recovered = True
                    break
        finally:
            image.close()
        if not recovered:
            unresolved.append(page["page"])
    document["unresolved_ocr_pages"] = unresolved
    document["orientation_attempts"] = attempts
    return document


def _matching_layout(path, ocr, enabled, literals):
    document = _layout(path, ocr)
    if ocr:
        _recover_orientation(path, document, enabled=enabled, literals=literals)
        matches, review = document_matches(document, enabled=enabled, literals=literals)
    else:
        matches = detect(document["text"], enabled=enabled, literals=literals)
        review = {
            "candidates": [],
            "unmatched_rule_ids": [],
            "short_or_repetitive_rule_ids": [],
            "unresolved_pages": [],
            "requires_confirmation": False,
            "scores_are_probabilities": False,
        }
    document["warnings"].extend(review_warnings(review))
    return document, matches, review


def clean_pdf(
    path,
    *,
    enabled=None,
    literals=None,
    ocr=False,
    rectangles=None,
    ocr_review_confirmed=False,
):
    from PIL import ImageDraw
    from pypdf import PdfReader, PdfWriter
    from localdesk.document_worker import render_page

    native = PdfReader(str(path), strict=False)
    if native.is_encrypted:
        raise InputError("Decrypt the PDF yourself before privacy processing.")
    if not 1 <= len(native.pages) <= 50:
        raise InputError("PDF redaction supports 1 to 50 complete pages.")
    document, matches, review = _matching_layout(path, ocr, enabled, literals)
    require_review(review, ocr_review_confirmed)
    if len(document["pages"]) != len(native.pages):
        raise InputError("Not all PDF pages were scanned. No clean copy was created.")
    selected = rectangles or []
    if not isinstance(selected, list) or len(selected) > 100:
        raise InputError("Use at most 100 manual page rectangles.")
    by_page = {}
    for item in selected:
        if not isinstance(item, dict):
            raise InputError(
                "Each PDF rectangle needs a page number and pixel coordinates."
            )
        page = integer(item.get("page"), 1, len(native.pages))
        r = item.get("rect")
        if (
            not isinstance(r, list)
            or len(r) != 4
            or any(isinstance(v, bool) or not isinstance(v, int) for v in r)
        ):
            raise InputError("PDF rectangle coordinates must be four integers.")
        by_page.setdefault(page, []).append(r)
    writer = PdfWriter()
    total_regions = 0
    for page in document["pages"]:
        image, _, _ = render_page(path, page["page"] - 1, scale=2)
        try:
            if image.width * image.height > MAX_PIXELS:
                raise InputError("A PDF page exceeds 20 million pixels.")
            regions = _regions(
                page["words"], matches, 2, image.width, image.height
            ) + by_page.get(page["page"], [])
            draw = ImageDraw.Draw(image)
            for x0, y0, x1, y1 in regions:
                if not 0 <= x0 < x1 <= image.width or not 0 <= y0 < y1 <= image.height:
                    raise InputError("A selected PDF rectangle is outside its page.")
                draw.rectangle((x0, y0, x1 - 1, y1 - 1), fill="black")
            total_regions += len(regions)
            fresh = io.BytesIO()
            image.save(fresh, format="PDF", resolution=144.0)
            new_page = PdfReader(io.BytesIO(fresh.getvalue()), strict=True)
            writer.add_page(new_page.pages[0])
        finally:
            image.close()
    # Writer receives only freshly encoded raster pages. No original page
    # objects, attachments, annotations, JavaScript, forms, or text survive.
    writer.add_metadata({"/Producer": "DataClean Room"})
    result = io.BytesIO()
    writer.write(result)
    data = result.getvalue()
    check = PdfReader(io.BytesIO(data), strict=True)
    if len(check.pages) != len(document["pages"]) or any(
        (p.extract_text() or "").strip() for p in check.pages
    ):
        raise InputError("The raster-only output failed verification.")
    return data, {
        "removed_matches": len(matches),
        "selected_regions": len(selected),
        "painted_regions": total_regions,
        "pages": len(check.pages),
        "output_format": "raster-only PDF",
        "source_objects_copied": False,
        "source_metadata_copied": False,
        "ocr_used": ocr,
        "ocr_review": review,
        "ocr_review_confirmed": ocr_review_confirmed,
        "orientation_attempts": document.get("orientation_attempts", 0),
        "human_review_required": True,
        "warnings": document.get("warnings", [])
        + [
            "The PDF is a new image-only copy. Searchable text, links, forms, annotations, and source attachments were removed. Annotations were excluded before rendering.",
            "Review every page. OCR and pattern checks can miss private text. Unselected faces and objects remain visible.",
        ],
    }


def normalize_pdf_geometry(raw):
    """Use a zero-origin, unrotated visible page box for scan and preview.

    Work on a copy. Raster export clips all content to this visible box.
    Drop annotations so their untransformed rectangles cannot confuse review.
    """
    from pypdf import PdfReader, PdfWriter, Transformation
    from pypdf.generic import RectangleObject, NameObject

    reader = PdfReader(io.BytesIO(raw), strict=False)
    if reader.is_encrypted:
        raise InputError("Decrypt the PDF yourself before privacy processing.")
    if not 1 <= len(reader.pages) <= 50:
        raise InputError("PDF redaction supports 1 to 50 complete pages.")
    writer = PdfWriter()
    for original in reader.pages:
        page = writer.add_page(original)
        if page.rotation:
            page.transfer_rotation_to_content()
        media, crop = page.mediabox, page.cropbox
        left, bottom = max(float(media.left), float(crop.left)), max(
            float(media.bottom), float(crop.bottom)
        )
        right, top = min(float(media.right), float(crop.right)), min(
            float(media.top), float(crop.top)
        )
        if (
            not all(math.isfinite(v) for v in (left, bottom, right, top))
            or right <= left
            or top <= bottom
        ):
            raise InputError("The PDF has an invalid visible page box.")
        if (right - left) * (top - bottom) * 4 > MAX_PIXELS:
            raise InputError("A PDF page exceeds 20 million rendered pixels.")
        page.add_transformation(Transformation().translate(tx=-left, ty=-bottom))
        box = RectangleObject((0, 0, right - left, top - bottom))
        for name in ("/MediaBox", "/CropBox", "/TrimBox", "/BleedBox", "/ArtBox"):
            page[NameObject(name)] = box
        if "/Annots" in page:
            del page["/Annots"]
    stream = io.BytesIO()
    writer.write(stream)
    result = stream.getvalue()
    if len(result) > MAX_FILE_BYTES:
        raise InputError("The normalized PDF exceeds 25 MiB.")
    return result


def process(raw, suffix, action, options):
    if len(raw) > MAX_FILE_BYTES:
        raise InputError("The source exceeds 25 MiB.")
    suffix = suffix.lower()
    if suffix not in OFFICE | {
        ".pdf",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".bmp",
        ".tif",
        ".tiff",
    }:
        raise InputError("Unsupported native privacy format.")
    enabled, literals = options.get("enabled"), options.get("literals")
    detect("", enabled=enabled, literals=literals)
    if action == "office":
        data, audit = clean_office(
            raw,
            suffix,
            enabled=enabled,
            literals=literals,
            mode=options.get("mode", "labels"),
        )
        return {
            "content": base64.b64encode(data).decode("ascii"),
            "audit": audit,
            "extension": suffix,
        }
    with tempfile.TemporaryDirectory(prefix="dataclean-parser-") as folder:
        path = Path(folder) / ("input" + suffix)
        path.write_bytes(normalize_pdf_geometry(raw) if suffix == ".pdf" else raw)
        if action == "layout":
            result, matches, review = _matching_layout(
                path, bool(options.get("ocr", False)), enabled, literals
            )
            return {
                "text": result["text"],
                "matches": public_matches(matches),
                "counts": dict(Counter(m["kind"] for m in matches)),
                "ocr_review": review,
                "automatic_regions": [
                    {"page": p["page"], "rect": r}
                    for p in result["pages"]
                    for r in _regions(
                        p["words"],
                        matches,
                        2 if suffix == ".pdf" else 1,
                        math.ceil(p["width"] * (2 if suffix == ".pdf" else 1)),
                        math.ceil(p["height"] * (2 if suffix == ".pdf" else 1)),
                    )
                ],
                "pages": [
                    {"page": p["page"], "width": p["width"], "height": p["height"]}
                    for p in result["pages"]
                ],
                "warnings": result.get("warnings", []),
                "method": result.get("method", "local layout"),
            }
        if action == "pdf":
            data, audit = clean_pdf(
                path,
                enabled=enabled,
                literals=literals,
                ocr=bool(options.get("ocr", False)),
                rectangles=options.get("rectangles"),
                ocr_review_confirmed=options.get("ocr_review_confirmed", False),
            )
            return {
                "content": base64.b64encode(data).decode("ascii"),
                "audit": audit,
                "extension": ".pdf",
            }
        if action == "preview":
            from localdesk.document_worker import render_page

            page = integer(options.get("page", 1), 1, 50)
            image, _, _ = render_page(path, page - 1, scale=2)
            try:
                stream = io.BytesIO()
                image.save(stream, format="PNG")
                return {
                    "content": base64.b64encode(stream.getvalue()).decode("ascii"),
                    "width": image.width,
                    "height": image.height,
                    "page": page,
                }
            finally:
                image.close()
        if action == "image":
            from .sanitizer import image_clean
            from PIL import Image, ImageOps

            with Image.open(io.BytesIO(raw)) as image:
                oriented = ImageOps.exif_transpose(image)
                width, height = oriented.size
            matches = []
            review = {"requires_confirmation": False}
            auto = []
            if options.get("ocr"):
                layout, matches, review = _matching_layout(
                    path, True, enabled, literals
                )
                require_review(review, options.get("ocr_review_confirmed", False))
                auto = _regions(layout["pages"][0]["words"], matches, 1, width, height)
            selected = options.get("rectangles") or []
            if len(auto) + len(selected) > 6000:
                raise InputError("The image has too many redaction regions.")
            # The image writer also enforces its 100-region limit.
            # Larger scans must be split; no region is silently discarded.
            regions = selected + auto
            data, notes = image_clean(raw, suffix, regions)
            return {
                "content": base64.b64encode(data).decode("ascii"),
                "extension": ".png",
                "audit": {
                    "removed_matches": len(matches),
                    "selected_regions": len(selected),
                    "automatic_regions": len(auto),
                    "ocr_review": review,
                    "ocr_review_confirmed": options.get("ocr_review_confirmed", False),
                    "source_metadata_copied": False,
                    "human_review_required": True,
                    "warnings": notes
                    + (layout["warnings"] if options.get("ocr") else []),
                },
            }
    raise InputError("This privacy operation is not supported.")
